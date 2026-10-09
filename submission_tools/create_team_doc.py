from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "submission" / "group_04_midproject" / "group_04_team_and_functionality.docx"

def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)

def add_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = value
        shade(c, "17324D")
        for r in c.paragraphs[0].runs:
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(9)
    for row in rows:
        cells = t.add_row().cells
        trpr = t.rows[-1]._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        for i, value in enumerate(row):
            cells[i].text = value
            for r in cells[i].paragraphs[0].runs:
                r.font.size = Pt(8.5)
            if len(t.rows) % 2 == 0:
                shade(cells[i], "EEF3F7")
    doc.add_paragraph()

def h(doc, text, level=1):
    p = doc.add_heading(text, level)
    p.paragraph_format.keep_with_next = True

def code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.75)
sec.right_margin = Inches(0.75)
styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
styles["Normal"].font.size = Pt(10)
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(24)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(23, 50, 77)
for name, size, color in (("Heading 1", 16, "17324D"), ("Heading 2", 12, "2854C5")):
    styles[name].font.name = "Aptos"
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].font.color.rgb = RGBColor.from_string(color)
header = sec.header.paragraphs[0]
header.text = "DI 31001 Socket Programming Networks | Group 4"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.runs[0].font.size = Pt(8)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("Group 4 technical explanation | 8 October 2026").font.size = Pt(8)

doc.add_paragraph("Team Responsibilities and Code Functionality", style="Title")
p = doc.add_paragraph()
r = p.add_run("IPv6 IRC Server and Command Bot | Mid-project Submission")
r.bold = True
r.font.name = "Aptos"
r.font.size = Pt(11)
r.font.color.rgb = RGBColor(40, 84, 197)
h(doc, "1 Team Responsibility")
doc.add_paragraph("The group assigns an equal contribution share of 20% to each member. Student A participated across every area and led the common protocol layer. B/C formed the server pair and D/E formed the Bot pair, with A participating across both. The detailed functions below identify each member's main implementation areas.")
add_table(doc, ["Member", "Responsibility", "Function and file evidence"], [
    ("A Kun Qian", "Common protocol and integration", "common/irc_message.py: IRCMessage.parse, serialize, to_bytes, to_bounded_bytes; common/stream.py: IRCStreamDecoder.feed; common/protocol.py: is_channel, valid_nickname, irc_casefold; common/numerics.py; protocol and integration review."),
    ("B Yibo Xu", "Server connection lifecycle", "server/server.py: serve_forever, send, send_many, broadcast_user_channels, disconnect, stop; server/client_session.py: start, run, send, close. IPv6 listener, sessions, serialized output and cleanup."),
    ("C Jingsong She", "Server protocol and state", "server/handlers.py: handle, error, registration, JOIN, PART, NAMES, PRIVMSG, PING, PONG, ISON and QUIT; server/state.py: identity, channel and routing operations."),
    ("D Tianxing Wei", "Bot transport and reliability", "bot/client.py: run, run_once, send, handle, close, stop. IPv6 registration, 001-triggered JOIN, PING/PONG, nickname recovery and reconnect behavior."),
    ("E Li Yuzong", "Bot state and commands", "bot/state.py: snapshots and JOIN/PART/QUIT/NICK events; bot/commands.py: channel_command and private_reply; bot/cli.py; bot/data/facts.txt."),
])
doc.add_paragraph("Shared work includes protocol interfaces among the common, server, and bot packages, regression tests, miniircd compatibility checks, and review of the final package. The official marking sheet assigns 20% to each member for every proposition.")

h(doc, "2 System Structure")
doc.add_paragraph("The common package contains protocol handling and no socket operations. The server package owns the IPv6 listener, client sessions, mutable user and channel state, and IRC command handlers. The bot package owns the client connection, local membership state, and command decisions. This separation makes parsing, response generation, and command selection testable without a live socket.")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.keep_with_next = True
p.add_run().add_picture(str(ROOT / ".codex-output" / "irc_architecture.png"), width=Inches(6.8))
p = doc.add_paragraph("Figure 1. System architecture and shared protocol dependencies.", style="Caption")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for image, caption in (("irc_package.png", "Figure 2. Package diagram."),
                       ("irc_class.png", "Figure 3. Class diagram.")):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(ROOT / ".codex-output" / image), width=Inches(6.8))
    p = doc.add_paragraph(caption, style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_table(doc, ["Layer", "Purpose", "Key structures"], [
    ("common", "IRC message representation and TCP framing", "IRCMessage, IRCStreamDecoder, numerics and RFC name validation"),
    ("server", "Registration, routing, channels, errors and cleanup", "IRCServer, ClientSession, ClientState, IRCState, IRCHandlers"),
    ("bot", "Connection, membership tracking and automatic replies", "BotClient, BotState, BotCommandProcessor, facts.txt"),
])

h(doc, "3 IRC Protocol Flow")
h(doc, "3.1 Message Model and TCP Framing", 2)
doc.add_paragraph("IRC is a CRLF-delimited text protocol carried over TCP. A recv call can contain half a command, one complete command, or several commands. IRCStreamDecoder.feed retains incomplete bytes, extracts complete CRLF lines, validates UTF-8, and rejects lines longer than 512 bytes, including CRLF. IRCMessage.parse separates an optional prefix, command, and parameters, while serialize and to_bytes produce the wire form.")
code(doc, ":alice!u@host PRIVMSG #hello :Hi Bob")
doc.add_paragraph("The logical parameters are #hello and Hi Bob. The colon before Hi Bob is a wire delimiter and is not stored in the logical text.")
h(doc, "3.2 Name Comparison and Registration", 2)
doc.add_paragraph("Nickname and channel comparisons use RFC 1459 casemapping. ASCII letters are compared case-insensitively, and [ ] \\ ^ map to { } | ~. This mapping is shared by collision checks, routing, Bot membership tracking, and !slap exclusions. A client sends NICK and USER; the server then sends numerics 001-004 after successful registration. The Bot waits for 001 before JOIN.")

h(doc, "4 Server Functionality and Error Handling")
add_table(doc, ["Requirement", "Implementation", "Evidence"], [
    ("Multiple clients", "AF_INET6 listener; one ClientSession per connection; IRCState protects clients, nicknames and channels with an RLock.", "test_integration.py and test_server_state.py"),
    ("Private messages", "PRIVMSG resolves one nickname and sends only to that client. Unknown targets return 401.", "IRCHandlers.cmd_privmsg"),
    ("Join and leave", "JOIN updates membership, broadcasts JOIN and sends 353/366. PART retains the connection. QUIT and disconnect clean all indexes.", "cmd_join, cmd_part, send_names and IRCState"),
    ("Channel messages", "Channel membership is required; messages go only to other members of the same channel.", "cmd_privmsg and IPv6 integration"),
    ("Errors", "Unknown commands, invalid or duplicate nicknames, missing parameters and invalid targets produce numerics. Malformed streams receive ERROR and close only that session.", "numerics.py and handler tests"),
    ("Keep-alive", "PING after 30 seconds without activity; removal after more than 60 seconds without command or PONG; scan every second.", "_check_idle_clients and timeout tests"),
    ("Wire length", "Relay messages are checked after adding the prefix. PART, QUIT and ERROR reasons are shortened at a UTF-8 boundary to remain within 512 bytes.", "IRCMessage.to_bounded_bytes and review tests"),
])

h(doc, "5 Bot Functionality")
h(doc, "5.1 Connection and Membership", 2)
doc.add_paragraph("BotClient creates an IPv6 socket, sends NICK and USER, answers PING with PONG, and joins after 001. A connection or stream failure closes the socket, waits three seconds, and reconnects with fresh state. Invalid or occupied nicknames are retried. The initial roster is collected from 353 replies and finalized at 366; JOIN, PART, QUIT, and NICK events then update it, including events that occur during a NAMES snapshot.")
h(doc, "5.2 Commands and Private Facts", 2)
add_table(doc, ["Command or event", "Behavior"], [
    ("Private PRIVMSG", "Replies with a random non-empty line from bot/data/facts.txt."),
    ("!hello", "Greets the channel sender."),
    ("!slap [user]", "Chooses an eligible member while excluding the Bot and sender when possible. Explicit self/Bot targets are rejected. An absent named target falls back to the sender, matching the brief."),
    ("!who", "Uses the new IRC ISON command and processes numeric 303 to report confirmed known members and a count."),
    ("!time", "Reports the Bot machine's local time with UTC offset."),
    ("NICK event", "Updates the member list and the Bot's own nickname when the server confirms a Bot rename."),
])

h(doc, "6 References")
references = [
    '[1] DI 31001, "Project Brief: Coursework 1 Socket Programming," University of Dundee, 2026.',
    '[2] DI 31001, "4. Marking Criteria," University of Dundee, 2026.',
    '[3] DI 31001, "Project Marking Sheet," University of Dundee, 2026.',
    '[4] C. Kalt, "Internet Relay Chat: Client Protocol," RFC 2812, Apr. 2000. [Online]. Available: https://www.rfc-editor.org/rfc/rfc2812',
    '[5] Python Software Foundation, "socket - Low-level networking interface," Python 3 Documentation. [Online]. Available: https://docs.python.org/3/library/socket.html. Accessed: Oct. 8, 2026.',
]
for reference in references:
    p = doc.add_paragraph(reference)
    for run in p.runs:
        run.font.size = Pt(9)
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.28)
    p.paragraph_format.space_after = Pt(1)
doc.core_properties.title = "Team Responsibilities and Code Functionality"
doc.core_properties.subject = "DI 31001 Socket Programming Networks Group 4"
doc.core_properties.author = "Group 4"
doc.save(OUT)
print(OUT)
